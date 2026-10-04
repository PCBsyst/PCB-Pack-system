"""Restyle built-in forms without replacing their fields, headers or watermarks."""
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
FILES = ["FGPC-008-01-application-review-kr.docx", "FGPC-012-01-decision-report-kr.docx", "FGPC-012-03-delivery-confirmation-en.docx"]

def fields(path):
    with ZipFile(path) as z:
        return sorted(re.findall(r"\{\{[^}]+\}\}", "".join(z.read(n).decode() for n in z.namelist() if n.endswith(".xml"))))

def tag(parent, name, **attrs):
    element = parent.find(qn("w:" + name))
    if element is None:
        element = OxmlElement("w:" + name)
        parent.append(element)
    for key, value in attrs.items():
        element.set(qn("w:" + key), str(value))
    return element

def paragraphs(table):
    # XML traversal avoids processing merged cells more than once.
    from docx.text.paragraph import Paragraph
    return [Paragraph(p, table) for p in table._tbl.iter(qn("w:p"))]

for name in FILES:
    path = ROOT / "templates" / name
    before = fields(path)
    doc = Document(path)
    original_text = [n.text for n in doc.element.body.iter(qn("w:t"))]
    for section in doc.sections:
        section.left_margin = section.right_margin = Mm(18)
        section.top_margin = section.bottom_margin = Mm(19)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for paragraph in doc.paragraphs + [p for t in doc.tables for p in paragraphs(t)]:
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            # Retain embedded pictures, field codes and typography emphasis.
            run.font.name = "Malgun Gothic"
            tag(run._element.get_or_add_rPr(), "rFonts", eastAsia="Malgun Gothic")
            if run.font.size is None or run.font.size.pt < 11:
                run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
    title = next((p for p in doc.paragraphs if p.text.strip()), None)
    if title is not None:
        if "Title" not in doc.styles:
            doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
        title.style = doc.styles["Title"]
        title.paragraph_format.space_after = Pt(16)
        title.paragraph_format.keep_with_next = True
        for run in title.runs:
            run.font.size = Pt(20)
            run.font.bold = True
    for table in doc.tables:
        borders = tag(table._tbl.tblPr, "tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            tag(borders, edge, val="single", sz="4", color="D9D9D9")
        margins = tag(table._tbl.tblPr, "tblCellMar")
        for side, width in [("top", 75), ("bottom", 75), ("left", 100), ("right", 100)]:
            tag(margins, side, w=width, type="dxa")
        for tc in table._tbl.iter(qn("w:tc")):
            tag(tc.get_or_add_tcPr(), "vAlign", val="center")
            shading = tc.tcPr.find(qn("w:shd"))
            if shading is not None and shading.get(qn("w:fill"), "auto") not in ["auto", "FFFFFF", "ffffff"]:
                shading.set(qn("w:fill"), "F2F4F7")
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            for height in list(tr_pr.findall(qn("w:trHeight"))):
                tr_pr.remove(height)
            tag(tr_pr, "cantSplit")
    assert original_text == [n.text for n in doc.element.body.iter(qn("w:t"))], name
    doc.save(path)
    assert before == fields(path), f"Template fields changed: {name}"
    print(f"Preserved all content and {len(before)} field occurrences: {name}")
