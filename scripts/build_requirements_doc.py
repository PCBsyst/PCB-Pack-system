from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "REQUIREMENTS.md"
OUTPUT = ROOT / "docs" / "Certification_Record_Management_System_요구사항_정의서.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 101, 115)


def set_font(run, name="Malgun Gothic", size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_font(run, size=9, color=GRAY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1


def add_title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("INTERNAL REQUIREMENTS SPECIFICATION"), size=10, bold=True, color=BLUE)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run("Certification Record\nManagement System"), size=27, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_font(subtitle.add_run("업무 요구사항 및 데이터 구조 정의서"), size=16, bold=True)
    meta = doc.add_table(rows=4, cols=2)
    meta.autofit = False
    meta.columns[0].width = Inches(1.5)
    meta.columns[1].width = Inches(5.0)
    rows = [("문서 상태", "현업 검토 초안"), ("작성 기준일", "2026-08-30"), ("적용 대상", "내부 자격인증 기록관리 업무"), ("핵심 원칙", "AI는 자격요건이나 승인 여부를 판단하지 않음")]
    for row, (label, value) in zip(meta.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], "E8EEF5")
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(1.5 if idx == 0 else 5.0)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_font(run, size=10.5, bold=(idx == 0))
    doc.add_page_break()


def add_inline_markup(paragraph, text, size=11):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(paragraph.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            set_font(paragraph.add_run(part), size=size)


def add_body_from_markdown(doc, markdown):
    lines = markdown.splitlines()
    in_code = False
    code_lines = []
    skip_front_title = True
    for line in lines:
        if line.startswith("```"):
            if in_code:
                table = doc.add_table(rows=1, cols=1)
                table.autofit = False
                table.columns[0].width = Inches(6.5)
                cell = table.cell(0, 0)
                shade_cell(cell, "F4F6F9")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                set_font(p.add_run("\n".join(code_lines)), name="Consolas", size=9, color=DARK_BLUE)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if skip_front_title and (line.startswith("# ") or line.startswith("## 업무 요구사항") or line.startswith("- 문서 상태:") or line.startswith("- 작성 기준일:") or line.startswith("- 적용 대상:") or line.startswith("- 원칙:" ) or not line.strip()):
            if line.startswith("## 1."):
                skip_front_title = False
            else:
                continue
        if line.startswith("## "):
            skip_front_title = False
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markup(p, line[3:], 16)
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markup(p, line[4:], 13)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline_markup(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markup(p, line[2:])
        elif line.strip():
            p = doc.add_paragraph()
            add_inline_markup(p, line.strip())


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("Certification Record Management System | 요구사항 정의서"), size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])
    add_title_page(doc)
    add_body_from_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
